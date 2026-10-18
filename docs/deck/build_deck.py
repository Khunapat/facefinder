"""Fill the AI Builder Cup template with FaceFinder's content."""
import copy
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

D = "/tmp/claude-0/-home-user/38c2b7c8-01f8-58c1-be00-aef99bc40a8b/scratchpad/deck/"
ASSETS = "/home/user/facefinder/design-canvas/assets/"
TEAM = sys.argv[1] if len(sys.argv) > 1 else "[TEAM NAME]"
VIDEO = sys.argv[2] if len(sys.argv) > 2 else "[ADD VIDEO LINK]"

INK = RGBColor(0x1F, 0x25, 0x40)
MUTED = RGBColor(0x5B, 0x61, 0x78)
ACCENT = RGBColor(0xD6, 0x33, 0x6C)
SOFT = RGBColor(0xFC, 0xEE, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation(D + "base.pptx")
S = list(prs.slides)


def box(slide, x, y, w, h, fill=SOFT, radius=True, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    if radius:
        shp.adjustments[0] = 0.08
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None, margin=0.0):
    """paras: list of (text or [(run, opts)], opts) ; opts: size, bold, color, align, after, bullet"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(margin))
    first = True
    for content, o in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = o.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(o.get("after", 4))
        runs = content if isinstance(content, list) else [(content, {})]
        for t, ro in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(ro.get("size", o.get("size", 12)))
            f.bold = ro.get("bold", o.get("bold", False))
            f.color.rgb = ro.get("color", o.get("color", INK))
        if o.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.18))))
            pPr.set("indent", str(-int(Inches(0.16))))
            from pptx.oxml.ns import qn
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for e in pPr.findall(qn(tag)):
                    pPr.remove(e)
            bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            pPr.append(bu)
    return tb


def picture(slide, path, x, y, w=None, h=None, name=None):
    im = Image.open(path)
    ratio = im.width / im.height
    if w and not h:
        h = w / ratio
    if h and not w:
        w = h * ratio
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        pic.name = name
    return pic, w, h


def append_after_label(slide, label, value, color=INK):
    """Add a run after the template's own label, copying its style."""
    for shp in slide.shapes:
        if not shp.has_text_frame:
            continue
        for p in shp.text_frame.paragraphs:
            if p.text.strip().startswith(label):
                src = p.runs[-1]
                new = copy.deepcopy(src._r)
                p._p.append(new)
                r = p.runs[-1]
                r.text = value
                r.font.color.rgb = color
                return
    raise SystemExit("label not found: " + label)


def keep_title_only(slide):
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip():
            ps = shp.text_frame.paragraphs
            for p in ps[1:]:
                p._p.getparent().remove(p._p)
            # the title paragraph ends in a line break; drop it
            for br in ps[0]._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}br"):
                ps[0]._p.remove(br)
            shp.height = Inches(0.61)


TODO = RGBColor(0xC0, 0x1E, 0x1E)

# ---- 1. team details
append_after_label(S[0], "Team name:", " " + TEAM, TODO if TEAM.startswith("[") else INK)
append_after_label(S[0], "Team leader name:", " Khunapat Prapassornchaikul")
append_after_label(S[0], "Problem Statement:", " Media, Content & Digital Experiences. FaceFinder turns one shared "
                   "Google Drive event album into one folder per person, with Gemini picking and sorting the best photos.")

# ---- 2. brief about the idea
s = S[1]
text(s, 0.45, 1.4, 4.75, 3.9, [
    ("The problem", {"bold": True, "color": ACCENT, "size": 13, "after": 2}),
    ("After an event, the photographer shares one Google Drive folder with hundreds of photos. Finding "
     "yourself, or sending friends their photos, means an hour of scrolling.", {"size": 13, "after": 12}),
    ("FaceFinder", {"bold": True, "color": ACCENT, "size": 13, "after": 2}),
    ("Scan the album once. A face model groups faces into people; Gemini judges and describes each photo. "
     "Then find a person, keep the best shots, sort by description, and copy one folder per person into "
     "your Drive. The album is never changed.", {"size": 13}),
], name="Brief text")
pic, w, h = picture(s, D + "img/everyone_live.png", 5.45, 1.5, w=4.1, name="Live app screenshot")
text(s, 5.45, 1.5 + h + 0.08, 4.1, 0.5, [("The live app on the demo album. The people are made up, created with Gemini.",
                                          {"size": 10, "color": MUTED})], name="Screenshot caption")

# ---- 3. opportunities
s = S[2]
keep_title_only(s)
cards = [
    ("How is it different?",
     "Event photos usually live in one shared Drive folder, with no face search and no culling. FaceFinder "
     "works on any shared Drive folder, for everyone who can open it, and gives back real Drive folders."),
    ("How does it solve the problem?",
     "One scan per album. Name a face or upload a selfie to get that person's photos. Gemini suggests "
     "keepers with a reason, and sorts by description. One click copies it all to your Drive."),
    ("USP",
     "Two AIs, each for its job: a face model decides who; Gemini decides how good and what is in it. "
     "Copy-only, no photo library kept, data deletable at any time. Works on a phone."),
]
for i, (q, a) in enumerate(cards):
    x = 0.45 + i * 3.1
    box(s, x, 1.45, 2.9, 3.55, name="Card %d" % (i + 1))
    text(s, x + 0.2, 1.65, 2.5, 3.2, [(q, {"bold": True, "color": ACCENT, "size": 15, "after": 10}), (a, {"size": 13})],
         name="Card %d text" % (i + 1))

# ---- 4. features
s = S[3]
feats = [
    ("cam-wave.png", "Everyone", "Every face grouped into people. Type a name to make that person's folder."),
    ("cam-search.png", "Find people + Best 10", "One selfie finds a person. Best 10 picks their best shots, each with Gemini's reason."),
    ("cam-photo.png", "Cull photos", "Keep, Review or Rejected, with reasons: blur, eyes closed, weaker shots of a burst."),
    ("cam-point.png", "Sort by description", "Describe a folder in plain words (\"throwing caps\"). Gemini fills it and suggests folders."),
    ("cam-celebrate.png", "Export to Drive", "One folder per person in your own Drive. Copies are made inside Google; never duplicated."),
    ("cam-stand.png", "Control and privacy", "Choose photos by day or by hand, stop and resume. See what is kept, delete it any time."),
]
for i, (icon, t, d) in enumerate(feats):
    col, row = i % 3, i // 3
    x, y = 0.45 + col * 3.1, 1.45 + row * 1.95
    box(s, x, y, 2.9, 1.8, name="Feature %d" % (i + 1))
    picture(s, ASSETS + icon, x + 0.15, y + 0.2, h=0.7, name="Feature %d icon" % (i + 1))
    text(s, x + 0.95, y + 0.18, 1.8, 1.5, [(t, {"bold": True, "color": ACCENT, "size": 13, "after": 4}), (d, {"size": 11})],
         name="Feature %d text" % (i + 1))

# ---- 5. process flow (a snake: 1-2-3 left to right, then down, 4-5-6 right to left)
s = S[4]
steps = [
    ("Sign in with Google", "Reads Drive with the user's own permission"),
    ("Paste a Drive folder link", "Private folder? Clear help to request access"),
    ("Choose photos", "All, by day, or by hand"),
    ("Scan once", "InsightFace detects and matches faces; Gemini writes one note per photo"),
    ("Work on the album", "Everyone, Find people, Cull, Sort"),
    ("Export to Drive", "Copies, one folder per person"),
]
pos = [(0.45, 1.45), (3.55, 1.45), (6.65, 1.45), (6.65, 3.3), (3.55, 3.3), (0.45, 3.3)]
for i, ((t, d), (x, y)) in enumerate(zip(steps, pos)):
    box(s, x, y, 2.9, 1.35, fill=SOFT, name="Step %d" % (i + 1))
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.15), Inches(y + 0.15), Inches(0.42), Inches(0.42))
    c.fill.solid(); c.fill.fore_color.rgb = ACCENT; c.line.fill.background(); c.shadow.inherit = False
    c.name = "Step %d number" % (i + 1)
    tf = c.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(i + 1)
    r.font.name = FONT; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE
    text(s, x + 0.7, y + 0.15, 2.1, 1.1, [(t, {"bold": True, "size": 13, "after": 3}), (d, {"size": 11, "color": MUTED})],
         name="Step %d text" % (i + 1))
arrows = [((3.35, 2.12), MSO_SHAPE.RIGHT_ARROW), ((6.45, 2.12), MSO_SHAPE.RIGHT_ARROW),
          ((6.45, 3.97), MSO_SHAPE.LEFT_ARROW), ((3.35, 3.97), MSO_SHAPE.LEFT_ARROW)]
for (x, y), kind in arrows:
    a = s.shapes.add_shape(kind, Inches(x), Inches(y - 0.1), Inches(0.2), Inches(0.2))
    a.fill.solid(); a.fill.fore_color.rgb = ACCENT; a.line.fill.background(); a.shadow.inherit = False
d = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(8.0), Inches(2.85), Inches(0.2), Inches(0.4))
d.fill.solid(); d.fill.fore_color.rgb = ACCENT; d.line.fill.background(); d.shadow.inherit = False
text(s, 0.45, 4.8, 9.1, 0.5, [([("Who uses it: ", {"bold": True}),
                                 ("the photographer culls and sorts; guests find themselves; friends share each person's folder.", {})],
                                {"size": 12, "color": MUTED})], name="Users line")

# ---- 6. wireframes / mock-ups: the phone layout of the live app
s = S[5]
shots = [("phone_start.png", "Start page"),
         ("phone_menu_c.png", "Album menu"),
         ("phone_sort.png", "Sort by description")]
for i, (f, cap) in enumerate(shots):
    x = 0.45 + i * 2.35
    pic, w, h = picture(s, D + "img/" + f, x, 1.35, h=3.6 if i != 1 else None, w=None if i != 1 else 1.8, name="Phone %d" % (i + 1))
    ln = pic.line; ln.color.rgb = RGBColor(0xD9, 0xDB, 0xE4); ln.width = Pt(0.75)
    text(s, x - 0.3, 1.35 + 3.62, w + 0.6, 0.4, [(cap, {"size": 11, "color": MUTED, "align": PP_ALIGN.CENTER})],
         name="Phone %d caption" % (i + 1))
text(s, 7.25, 1.5, 2.3, 3.2, [("Phone first", {"bold": True, "color": ACCENT, "size": 13, "after": 4}),
                              ("Works from 320 px wide to desktop, with large touch targets and a bottom action bar. Our own pink design and camera mascot.", {"size": 12})],
     name="Phone note")

# ---- 7. architecture
s = S[6]
picture(s, D + "img/architecture_c.png", 0.55, 1.35, w=8.9, name="Architecture diagram")

# ---- 8. technologies
s = S[7]
cols = [
    ("Google Cloud", [
        ("Cloud Run", "runs the app and the face model"),
        ("Firestore", "scan results, Gemini notes, names"),
        ("Firebase Hosting", "the public address"),
        ("Vertex AI", "Gemini, no API key needed"),
        ("Drive API, Google sign-in", "read albums, copy photos"),
        ("Secret Manager", "sign-in secrets"),
    ]),
    ("AI models", [
        ("Gemini 3.5 Flash", "photo notes, sorting, folder suggestions"),
        ("InsightFace buffalo_l", "face detection and matching, on CPU"),
        ("OpenCV", "images, burst fingerprints"),
    ]),
    ("App and tools", [
        ("Python, FastAPI", "server and background jobs"),
        ("HTML, CSS, JavaScript", "one-file front end"),
        ("Docker", "one container"),
        ("Tests", "90 Python checks, browser tests"),
        ("Antigravity, Gemini", "testing, demo photos of made-up people"),
    ]),
]
for i, (head, items) in enumerate(cols):
    x = 0.45 + i * 3.1
    box(s, x, 1.45, 2.9, 3.85, name="Tech card %d" % (i + 1))
    paras = [(head, {"bold": True, "color": ACCENT, "size": 14, "after": 8})]
    for k, v in items:
        paras.append(([(k, {"bold": True}), (": " + v, {})], {"size": 11, "after": 5}))
    text(s, x + 0.2, 1.6, 2.5, 3.6, paras, name="Tech card %d text" % (i + 1))

# ---- 9. cost
s = S[8]
text(s, 0.45, 1.45, 5.6, 3.9, [
    ("What it costs to run", {"bold": True, "color": ACCENT, "size": 14, "after": 8}),
    ([("Cloud Run: ", {"bold": True}), ("one 4-vCPU instance; it scales to zero when idle.", {})], {"bullet": True, "after": 8, "size": 13}),
    ([("Gemini: ", {"bold": True}), ("one image call per photo. Cached notes avoid repeating it for unchanged photos; "
                                    "Sort sends text only.", {})], {"bullet": True, "after": 8, "size": 13}),
    ([("After the scan: ", {"bold": True}), ("Find people, Cull and Best 10 reuse saved notes, with no new Gemini calls.", {})], {"bullet": True, "after": 8, "size": 13}),
    ([("Still charged: ", {"bold": True}), ("compute, Firestore storage and reads, as usual. No measured per-album cost is claimed.", {})], {"bullet": True, "after": 8, "size": 13}),
], name="Cost text")
box(s, 6.4, 1.6, 3.15, 3.4, name="Cost callout")
text(s, 6.6, 1.8, 2.75, 3.0, [
    ("Per 1,000 photos", {"bold": True, "color": ACCENT, "size": 13, "after": 6}),
    ("1,000", {"bold": True, "size": 30, "after": 0}),
    ("Gemini image calls, once", {"size": 11, "color": MUTED, "after": 10}),
    ("~25", {"bold": True, "size": 30, "after": 0}),
    ("text calls for each sort", {"size": 11, "color": MUTED, "after": 10}),
    ("0", {"bold": True, "size": 30, "after": 0}),
    ("new Gemini image calls to rescan unchanged photos", {"size": 11, "color": MUTED}),
], name="Cost numbers")

# ---- 10. snapshots
s = S[9]
pic, w, h = picture(s, D + "img/everyone_live.png", 0.3, 1.3, w=6.0, name="Snapshot Everyone")
text(s, 0.3, 1.3 + h + 0.06, 6.0, 0.5, [("Everyone on the live app: name a face to make a folder. Right: scanning, and Sort by "
                                           "description.", {"size": 10, "color": MUTED})], name="Snapshot caption")
picture(s, D + "img/scan_c.png", 6.45, 1.3, w=3.3, name="Snapshot scan")
picture(s, D + "img/sort_c.png", 6.45, 2.85, w=3.3, name="Snapshot sort")
for shp in list(s.shapes)[-4:]:
    if shp.shape_type == 13:
        shp.line.color.rgb = RGBColor(0xD9, 0xDB, 0xE4); shp.line.width = Pt(0.75)

# ---- 11. performance
s = S[10]
stats = [
    ("0.75–0.98", "face match score for the same person, against −0.05 for a different person"),
    ("40 of 40", "small-face stress test: 40 faces 18 px wide in a synthetic grid"),
    ("1 min 36 s", "one observed live run: 37 photos scanned with Gemini notes (4 vCPU)"),
    ("28 s", "one observed live run: 31 copies exported into 5 Drive folders"),
    ("16 s", "one observed live run: upload a selfie, all 3 of that person's photos found"),
    ("5.2 s", "one observed live run: Stop ends a scan"),
]
for i, (big, label) in enumerate(stats):
    col, row = i % 3, i // 3
    x, y = 0.45 + col * 3.1, 1.45 + row * 1.8
    box(s, x, y, 2.9, 1.6, name="Stat %d" % (i + 1))
    text(s, x + 0.2, y + 0.15, 2.5, 1.35, [(big, {"bold": True, "color": ACCENT, "size": 26, "after": 2}),
                                           (label, {"size": 11})], name="Stat %d text" % (i + 1))
text(s, 0.45, 5.1, 9.1, 0.35, [("Scores and the stress test: the repository's automatic tests. Timings: single observed runs on the live app, 17–18 October 2026.",
                                {"size": 10, "color": MUTED})], name="Stats source")

# ---- 12. additional details / future
s = S[11]
pic, w, h = picture(s, D + "img/privacy.png", 0.45, 1.45, w=4.4, name="Privacy screenshot")
pic.line.color.rgb = RGBColor(0xD9, 0xDB, 0xE4); pic.line.width = Pt(0.75)
text(s, 0.45, 1.45 + h + 0.08, 4.4, 0.8, [("Trust today: no photo library kept, a plain list of what is kept, a public privacy policy, "
                                           "and \"Delete this album's data\".", {"size": 10, "color": MUTED})],
     name="Privacy caption")
text(s, 5.2, 1.45, 4.4, 3.9, [
    ("Next steps", {"bold": True, "color": ACCENT, "size": 14, "after": 8}),
    ([("Scale: ", {"bold": True}), ("persist job progress in Firestore, then run scans on Cloud Run jobs or Cloud Tasks.", {})], {"bullet": True, "after": 8, "size": 13}),
    ([("Guest mode: ", {"bold": True}), ("a QR code at the event; each guest finds their own photos.", {})], {"bullet": True, "after": 8, "size": 13}),
    ([("Commercial use: ", {"bold": True}), ("a commercially licensed face model.", {})], {"bullet": True, "after": 8, "size": 13}),
    ([("Google verification: ", {"bold": True}), ("no more \"unverified app\" warning.", {})], {"bullet": True, "size": 13}),
], name="Future text")

# ---- 13. links
s = S[12]
for label, val, col in (("GitHub Public Repository", "https://github.com/Khunapat/facefinder", INK),
                        ("Demo Video Link", VIDEO, TODO if VIDEO.startswith("[") else INK),
                        ("Final Product Link", "https://facefinder-508213.web.app", INK)):
    append_after_label(s, label, ": ")
    append_after_label(s, label, val, col)
for shp in s.shapes:
    if shp.has_text_frame and "GitHub" in shp.text_frame.text:
        shp.height = Inches(2.0)
        for p in shp.text_frame.paragraphs:
            for r in p.runs:
                if "Demo Video Link (3 Minutes)" in r.text:
                    r.text = r.text.replace("(3 Minutes)", "(3 to 4 minutes)")
text(s, 0.45, 2.45, 9.1, 1.5, [
    ("For judges", {"bold": True, "color": ACCENT, "size": 13, "after": 4}),
    ("Open the app and sign in with Google. Google may show \"Google hasn't verified this app\": press Advanced, then "
     "Go to facefinder. Open the Sample album card (made-up people created with Gemini) and try Everyone, Find people, "
     "Sort and Cull photos.", {"size": 12}),
], name="Judge note")

LINKS = {"https://github.com/Khunapat/facefinder": "https://github.com/Khunapat/facefinder",
         "https://facefinder-508213.web.app": "https://facefinder-508213.web.app"}
if VIDEO.startswith("http"):
    LINKS[VIDEO] = VIDEO
for shp in S[12].shapes:
    if not shp.has_text_frame:
        continue
    for p in shp.text_frame.paragraphs:
        for r in p.runs:
            for url in LINKS:
                if r.text.strip() == url:
                    r.hyperlink.address = url
prs.save(D + "FaceFinder_deck.pptx")
print("saved")
